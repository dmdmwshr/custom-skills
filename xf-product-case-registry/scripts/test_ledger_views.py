import hashlib
import json
from datetime import date
from pathlib import Path

import pytest

from scripts import ledger_views as views
from scripts import source_intake as source
from scripts import workspace_state as ws

A = "99999999T209900001"
B = "99999999T209900002"
NOW = "2099-02-20T10:00:00+08:00"


@pytest.fixture
def layout(tmp_path):
    return ws.ensure_workspace_layout(ws.BusinessLayout.from_root(tmp_path / "work"))


def write(path, value):
    path.parent.mkdir(parents=True, exist_ok=True)
    path.write_text(json.dumps(value), encoding="utf-8")


def add_case(layout, project, result="UNQUALIFIED", archived=False, deep=False):
    work = layout.work_case_dir(project)
    generation = "20990220T020000Z-aaaaaaaaaaaa"
    if archived:
        work = layout.root / "工作区" / "历史工作区" / f"{project}-{generation}"
    manifest = {
        "case": {"projectNo": project},
        "packageSha256": "sha256:" + "a" * 64,
        "initialInspection": {
            "inspectionDate": "2098-12-10",
            "products": [{"result": result, "method": "ONSITE"}],
        },
        "files": [{"clientRef": "file-a", "sha256": "sha256:" + "b" * 64}],
    }
    write(work / "manifest.json", manifest)
    digest = "sha256:" + hashlib.sha256((work / "manifest.json").read_bytes()).hexdigest()
    state = {
        "stateVersion": 6,
        "projectNo": project,
        "manifestSha256": digest,
        "caseId": "fixture-case",
        "origin": "https://fixture.example",
        "status": "VERIFIED",
        "finalizeSummary": {"created": True, "conflictCount": 0, "skippedCount": 0},
        "filesProjection": [{"clientRef": "file-a"}],
        "uploadedFileRefs": ["file-a"],
    }
    write(work / "upload-state.json", state)
    archive = {}
    if archived:
        receipt_path = layout.verification_records / f"{project}-receipt.json"
        write(
            receipt_path,
            {
                "recordVersion": 1,
                "projectNo": project,
                "manifestSha256": digest,
                "archivedWorkspace": str(work),
            },
        )
        archive = {"workspacePath": str(work), "verificationRecord": str(receipt_path)}
    if deep:
        write(
            work / "content-verification.json",
            {
                "schemaVersion": "ContentVerificationV1",
                "manifestSha256": digest,
                "completedAt": NOW,
                "caseId": state["caseId"],
                "origin": state["origin"],
                "files": {
                    "file-a": {
                        "sha256": "sha256:" + "b" * 64,
                        "fileId": "fixture-file",
                        "verifiedAgainstManifestSha256": digest,
                        "verifiedAt": NOW,
                        "contentGeneration": 1,
                        "sizeBytes": "1",
                    }
                },
            },
        )
    ws.upsert_case(
        layout,
        project,
        state="COMPLETED" if archived else "UPLOADED_AWAITING_NAS",
        source={
            "status": "PACKAGE_READY",
            "rwid": project,
            "batchId": "first",
            "projectIdentitySource": "DETAIL",
        },
        upload={"status": "VERIFIED"},
        nasVerification={"status": "VERIFIED"},
        archive=archive,
    )


def test_all_and_unqualified_share_identity_and_include_archived_file_counts(layout):
    add_case(layout, A, archived=True, deep=True)
    add_case(layout, B, result="QUALIFIED")
    ws.upsert_case(layout, A, source={"batchId": "second"})
    full = views.ledger_view(layout)
    subset = views.ledger_view(layout, view="unqualified")
    assert full["counts"]["cases"] == 2
    assert full["counts"]["receivedFiles"] == 2
    assert subset["counts"]["cases"] == 1
    assert subset["cases"][0]["projectNo"] == A
    assert subset["cases"][0]["complete"] is True
    assert subset["cases"][0]["statisticsYear"] == 2098
    assert ws.load_waterline(layout)["cases"][A]["source"]["batchIds"] == ["first", "second"]


def test_legacy_verified_is_not_deep_content_proof(layout):
    add_case(layout, A, archived=True)
    row = views.ledger_view(layout)["cases"][0]
    assert row["historicalComplete"] is True
    assert row["stages"]["bodyVerified"] is False
    assert row["complete"] is False
    assert views.scan_plan(layout)["pending"] == []  # no automatic historical re-download


def test_workflow_mutation_does_not_refresh_source_observation(layout):
    ws.upsert_case(layout, A, source={"lastObservedAt": NOW})
    ws.upsert_case(layout, A, upload={"status": "UPLOADING"})
    record = ws.load_waterline(layout)["cases"][A]
    assert record["lastSeenAt"] == NOW
    assert record["source"]["lastObservedAt"] == NOW
    assert record["workflowUpdatedAt"]


def test_rolling_calendar_window_handles_year_and_short_months():
    assert views.recent_window(date(2026, 9, 20)) == {
        "startDate": "2026-06-20",
        "endDate": "2026-09-20",
    }
    assert views.recent_window(date(2024, 5, 31))["startDate"] == "2024-02-29"
    assert views.recent_window(date(2026, 2, 20))["startDate"] == "2025-11-20"


def test_reinspection_result_and_recheck_are_independent():
    manifest = {
        "initialInspection": {
            "products": [
                {
                    "method": "SAMPLING",
                    "result": "UNQUALIFIED",
                    "reinspectionApplied": "YES",
                    "reinspectionResult": "QUALIFIED",
                }
            ]
        },
        "recheckInspection": {"products": [{"result": "UNQUALIFIED"}]},
    }
    assert views.qualification(manifest)["initialResult"] == "QUALIFIED"
    manifest["initialInspection"]["products"][0]["reinspectionResult"] = "PENDING"
    assert views.qualification(manifest)["initialResult"] == "UNKNOWN"
    manifest["initialInspection"]["products"].append({"method": "ONSITE", "result": "UNQUALIFIED"})
    assert views.qualification(manifest)["initialResult"] == "UNQUALIFIED"


def test_source_initial_onsite_failure_survives_recheck_and_sampling():
    products = [
        {
            "页面记录": {"SFCPFC": "0", "JCRQ": "2099-05-26 14:30:00", "SFHG": "1"},
            "检查结果": "",
            "检查结果标记": ["el-icon-error"],
            "产品质量现场检查情况": "不合格",
        },
        {
            "页面记录": {"SFCPFC": "0"},
            "检查结果": "[检]",
            "检查结果标记": ["el-icon-success"],
            "产品质量现场检查情况": "未发现不合格现象",
        },
        {
            "页面记录": {"SFCPFC": "1", "SFHG": "0"},
            "检查结果": "[复查]",
            "检查结果标记": ["el-icon-success"],
            "产品质量现场检查情况": "未发现不合格现象",
        },
    ]
    result = views.source_qualification({"检查产品信息": products}, [])
    assert result["initialResult"] == "UNQUALIFIED"
    assert result["statisticsYear"] == 2099
    assert views.source_qualification({"检查产品信息": products[1:]}, []) is None


@pytest.mark.parametrize(
    "change",
    [
        {"检查结果": "[检]"},
        {"检查结果": "[复检]"},
        {"检查结果": "[复查]"},
        {"页面记录": {"SFCPFC": "1"}},
        {"页面记录": {}},
        {"检查结果标记": ["el-icon-success"]},
        {"检查结果标记": []},
        {"产品质量现场检查情况": "未发现不合格现象"},
    ],
)
def test_source_notice_or_ambiguous_product_is_not_final_unqualified(change):
    product = {
        "页面记录": {"SFCPFC": "0", "SFHG": "1"},
        "检查结果": "",
        "检查结果标记": ["el-icon-error"],
        "产品质量现场检查情况": "不合格",
    }
    product.update(change)
    fields = {
        "检查产品信息": [product],
        "检查情况": "已结案[复查不合格]",
        "文书目录": ["责令限期改正通知书"],
    }
    assert views.source_qualification(fields, []) is None


def test_incremental_contract_all_types_and_old_pending_survives(layout):
    filters = {**views.scan_filters(date(2099, 2, 20)), "queryEvidencePath": "query.json"}
    clean = source._default_filters(NOW, filters)
    assert clean["startDate"] == "2098-11-20"
    assert clean["documentType"] == "ALL"
    with pytest.raises(source.SourceIntakeError):
        source._default_filters(NOW, {**filters, "documentType": "责令限期改正通知书"})
    add_case(layout, A)
    assert views.scan_plan(layout)["pending"][0]["projectNo"] == A


def test_known_documents_skip_without_entering_detail_and_new_document_does_not(layout):
    add_case(layout, A, archived=True)
    appearance = {"文书名称": "检查记录", "创建日期": "2099-01-20"}
    ws.upsert_case(
        layout,
        A,
        source={
            "rwid": "fixture-rwid",
            "observationsByRwid": {
                "fixture-rwid": {"documentFingerprints": [source._fingerprint(appearance)]}
            },
        },
    )
    state = {"filters": {"selectionMode": "RECENT_DOCUMENT_ACTIVITY"}, "batchId": "recent"}
    records = {
        "fixture-rwid": {"rwid": "fixture-rwid", "fields": {}, "sourceAppearances": [appearance]}
    }
    source._prepare_incremental_queue(layout, state, records, NOW)
    assert state["actionRwids"] == []
    assert records["fixture-rwid"]["skippedAsUnchanged"] is True
    records = {
        "fixture-rwid": {
            "rwid": "fixture-rwid",
            "fields": {},
            "sourceAppearances": [appearance, {"文书名称": "复查记录", "创建日期": "2099-02-20"}],
        }
    }
    source._prepare_incremental_queue(layout, state, records, NOW)
    assert state["actionRwids"] == ["fixture-rwid"]


def test_bad_archive_reference_cannot_supply_registration_evidence(layout):
    ws.upsert_case(layout, A, state="COMPLETED", archive={"workspacePath": str(layout.root.parent)})
    row = views.ledger_view(layout)["cases"][0]
    assert row["stages"]["systemRegistered"] is False
    assert row["issues"] == ["ARCHIVE_REFERENCE_INVALID"]


def test_system_snapshot_values_preserve_effective_result_and_initial_year(layout):
    snapshot = {
        "inspections": [
            {
                "stage": "INITIAL_CHECK",
                "fields": {"inspectionDate": {"value": "2098-12-10", "source": "MANUAL"}},
                "products": [
                    {
                        "fields": {
                            "method": {"value": "SAMPLING"},
                            "result": {"value": "UNQUALIFIED"},
                            "reinspectionApplied": {"value": "YES"},
                            "reinspectionResult": {"value": "QUALIFIED"},
                        }
                    }
                ],
            },
            {"stage": "RECHECK", "products": [{"fields": {"result": {"value": "UNQUALIFIED"}}}]},
        ]
    }
    classification = views.snapshot_qualification(snapshot)
    assert classification["initialResult"] == "QUALIFIED"
    assert classification["statisticsYear"] == 2098
    ws.upsert_case(
        layout,
        A,
        systemObservation={
            "exists": True,
            "caseId": "known-system-case",
            "snapshotDigest": "sha256:" + "a" * 64,
            "qualification": classification,
        },
    )
    row = views.ledger_view(layout)["cases"][0]
    assert row["initialResult"] == "QUALIFIED"
    assert row["stages"]["systemRegistered"] is True
    assert row["stages"]["materialsCollected"] is False


def test_recent_activity_keeps_archive_and_only_queues_new_document_evidence(layout):
    add_case(layout, A, result="QUALIFIED", archived=True, deep=True)
    old_document = {"文书名称": "检查记录", "创建日期": "2098-12-10"}
    ws.upsert_case(
        layout,
        A,
        source={
            "rwid": "fixture-rwid",
            "observationsByRwid": {
                "fixture-rwid": {
                    "documentFingerprints": [source.document_fingerprint(old_document)]
                }
            },
        },
    )
    archive = ws.load_waterline(layout)["cases"][A]["archive"].copy()
    work = Path(archive["workspacePath"])
    original = {p.name: p.read_bytes() for p in work.iterdir() if p.is_file()}
    filters = {**views.scan_filters(date(2099, 2, 20)), "queryEvidencePath": "synthetic-query.json"}
    source.begin_capture(
        layout, filters, origin="https://source.example/", now=NOW, batch_id="new-doc"
    )
    row = {
        "RWID": "fixture-rwid",
        "单位名称": "测试单位",
        "文书名称": "抽样复检报告",
        "创建日期": "2099-02-20",
    }
    for round_no in (1, 2):
        source.add_page(layout, "new-doc", 1, [row], 1, 1, round_no=round_no, observed_at=NOW)
        source.finalize_capture(layout, "new-doc", now=NOW)
    detail = {
        "projectNo": A,
        "unitName": "测试单位",
        "initialInspection": {
            "inspectionDate": "2098-12-10",
            "products": [{"method": "ONSITE", "result": "UNQUALIFIED"}],
        },
    }
    source.add_detail(layout, "new-doc", "fixture-rwid", detail, captured_at=NOW)
    case = ws.load_waterline(layout)["cases"][A]
    assert case["archive"] == archive
    assert original == {p.name: p.read_bytes() for p in work.iterdir() if p.is_file()}
    assert not layout.work_case_dir(A).exists()
    planned = views.scan_plan(layout)["pending"]
    assert planned[0]["initialResult"] == "UNQUALIFIED"
    assert planned[0]["changeEvidence"][0]["newDocumentCount"] == 1
    evidence_path = layout.root / planned[0]["changeEvidence"][0]["evidencePath"]
    before = evidence_path.read_bytes()
    source.add_detail(layout, "new-doc", "fixture-rwid", detail, captured_at=NOW)
    assert evidence_path.read_bytes() == before
    current = views.ledger_view(layout)
    assert current["counts"]["cases"] == 1
    assert current["cases"][0]["complete"] is False
    assert current["cases"][0]["stages"]["bodyVerified"] is True


def test_excel_views_share_one_snapshot_and_keep_business_stages_independent(layout):
    from openpyxl import load_workbook

    add_case(layout, A, archived=True, deep=True)
    ws.upsert_case(
        layout,
        B,
        systemObservation={
            "exists": True,
            "caseId": "case-only-in-system",
            "snapshotDigest": "sha256:" + "b" * 64,
            "qualification": {
                "initialResult": "UNQUALIFIED",
                "initialInspectionDate": "2098-11-20",
                "statisticsYear": 2098,
                "reinspectionPending": False,
            },
        },
    )
    before = layout.waterline_json.read_bytes()
    target = ws.export_waterline_xlsx(layout)
    assert layout.waterline_json.read_bytes() == before
    with __import__("contextlib").closing(load_workbook(target)) as book:
        assert book.sheetnames[:4] == ["所有案卷", "不合格案卷", "待确认案卷", "未完成案卷"]
        assert book["所有案卷"].max_row == book["不合格案卷"].max_row == 3
        assert book["待确认案卷"].max_row == 1
        assert book["未完成案卷"]["A2"].value == B
        assert book["未完成案卷"]["G2"].value == "未确认"
        assert book["未完成案卷"]["H2"].value == "是"


def native_qualified_fields():
    return {
        "检查情况": "已结案[合格]",
        "预定检查日期": "2098-11-20",
        "创建日期": "2098-11-21",
        "检查产品信息": [
            {
                "产品质量现场检查情况": "未发现不合格现象",
                "市场准入检查情况": "未发现不合格现象",
                "检查结果": "",
            }
        ],
    }


def test_native_source_classification_requires_exact_initial_product_evidence():
    fields = native_qualified_fields()
    observed = views.source_qualification(fields, ["INITIAL"])
    assert observed["initialResult"] == "QUALIFIED"
    assert observed["statisticsYear"] is None  # planned/created dates are not actual inspection
    assert views.source_qualification(fields, ["INITIAL", "RECHECK"]) is None
    assert views.source_qualification(fields, []) is None
    fields["检查产品信息"][0]["产品质量现场检查情况"] = ""
    assert views.source_qualification(fields, ["INITIAL"]) is None


def test_saved_source_classification_is_read_only_until_applied_and_reuses_index(layout):
    ws.upsert_case(layout, A, source={"projectIdentitySource": "DETAIL"})
    path = layout.pending_case_dir(A) / "source-evidence.json"
    write(
        path,
        {
            "projectNo": A,
            "updatedAt": NOW,
            "records": {
                "rwid": {
                    "projectNo": A,
                    "fields": native_qualified_fields(),
                    "projectInspectionStages": ["INITIAL"],
                },
            },
        },
    )
    before = layout.waterline_json.read_bytes()
    assert views.classify_saved_details(layout)[0]["initialResult"] == "QUALIFIED"
    assert layout.waterline_json.read_bytes() == before
    assert views.index_history(layout, apply=True)["classificationChanges"][0]["applied"]
    assert views.classify_saved_details(layout, apply=True) == []
    row = views.ledger_view(layout)["cases"][0]
    assert row["initialResult"] == "QUALIFIED"
    assert row["stages"]["materialsCollected"] is False
    assert row["stages"]["systemRegistered"] is False
    assert views.scan_plan(layout)["pending"] == []
    path.write_text("{}", encoding="utf-8")
    assert views.ledger_view(layout)["cases"][0]["initialResult"] == "UNKNOWN"


def test_saved_source_classification_does_not_discard_ambiguous_or_foreign_evidence(layout):
    ws.upsert_case(layout, A, source={"projectIdentitySource": "DETAIL"})
    proof = {
        "projectNo": A,
        "records": {
            "first": {
                "projectNo": A,
                "fields": native_qualified_fields(),
                "projectInspectionStages": ["INITIAL"],
            },
            "second": {
                "projectNo": A,
                "fields": native_qualified_fields(),
                "projectInspectionStages": ["RECHECK"],
            },
        },
    }
    path = layout.pending_case_dir(A) / "source-evidence.json"
    write(path, proof)
    assert views.classify_saved_details(layout, apply=True) == []
    proof["records"].pop("second")
    proof["projectNo"] = B
    write(path, proof)
    assert views.classify_saved_details(layout, apply=True) == []


def test_current_system_reinspection_overrides_old_manifest_result(layout):
    add_case(layout, A, archived=True, deep=True)
    ws.upsert_case(
        layout,
        A,
        systemObservation={
            "exists": True,
            "caseId": "fixture-case",
            "snapshotDigest": "sha256:" + "c" * 64,
            "qualification": {
                "initialResult": "QUALIFIED",
                "initialInspectionDate": "2098-12-10",
                "statisticsYear": 2098,
                "reinspectionPending": False,
            },
        },
    )
    assert views.ledger_view(layout)["cases"][0]["initialResult"] == "QUALIFIED"
    assert views.ledger_view(layout, view="unqualified")["counts"]["cases"] == 0


def test_changed_system_generation_reopens_only_that_case_without_touching_archive(layout):
    add_case(layout, A, archived=True, deep=True)
    add_case(layout, B, archived=True, deep=True)
    archive = ws.load_waterline(layout)["cases"][A]["archive"].copy()
    ws.upsert_case(
        layout,
        A,
        systemObservation={
            "exists": True,
            "caseId": "fixture-case",
            "snapshotDigest": "sha256:" + "c" * 64,
            "fileSnapshot": {
                "fixture-file": {
                    "sha256": "sha256:" + "b" * 64,
                    "sizeBytes": 1,
                    "contentGeneration": 2,
                }
            },
        },
    )
    current = views.ledger_view(layout)
    assert current["cases"][0]["systemChanged"] is True
    assert current["cases"][0]["stages"]["bodyVerified"] is False
    assert current["cases"][0]["historicalComplete"] is True
    assert [item["projectNo"] for item in views.scan_plan(layout)["pending"]] == [A]
    assert ws.load_waterline(layout)["cases"][A]["archive"] == archive


def test_confirmed_missing_system_case_is_not_reported_registered_from_old_state(layout):
    add_case(layout, A, archived=True, deep=True)
    ws.upsert_case(layout, A, systemObservation={"exists": False, "observedAt": NOW})
    row = views.ledger_view(layout)["cases"][0]
    assert row["stages"]["systemRegistered"] is False
    assert row["complete"] is False
    assert row["activePending"] is True


def test_unchanged_current_file_identity_reuses_receipt(layout):
    add_case(layout, A, archived=True, deep=True)
    ws.upsert_case(
        layout,
        A,
        systemObservation={
            "exists": True,
            "caseId": "fixture-case",
            "snapshotDigest": "sha256:" + "c" * 64,
            "fileSnapshot": {
                "fixture-file": {
                    "sha256": "sha256:" + "b" * 64,
                    "sizeBytes": "1",
                    "contentGeneration": 1,
                }
            },
        },
    )
    row = views.ledger_view(layout)["cases"][0]
    assert row["stages"]["bodyVerified"] is True
    assert row["complete"] is True
    assert views.scan_plan(layout)["pending"] == []


def test_daily_report_only_expands_pending_or_anomalous_cases(layout):
    add_case(layout, A, archived=True, deep=True)
    add_case(layout, B)
    snapshot = views.ledger_view(layout)
    daily = views.render_ledger_report(snapshot)
    assert A not in daily and B in daily
    assert "本视图 2 案" in daily
    assert A in views.render_ledger_report(snapshot, include_all=True)
