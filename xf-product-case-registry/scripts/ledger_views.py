"""Read-only views over the one project-number keyed case waterline."""

from __future__ import annotations

import calendar
import hashlib
from collections import Counter
from datetime import date, datetime
from pathlib import Path
from typing import Any
from zoneinfo import ZoneInfo

if __package__:
    from . import workspace_state as ws
else:
    import workspace_state as ws


def recent_window(today: date | None = None) -> dict[str, str]:
    today = today or datetime.now(ZoneInfo("Asia/Shanghai")).date()
    month_index = today.year * 12 + today.month - 1 - 3
    year, month = divmod(month_index, 12)
    month += 1
    start = date(year, month, min(today.day, calendar.monthrange(year, month)[1]))
    return {"startDate": start.isoformat(), "endDate": today.isoformat()}


def scan_filters(today: date | None = None) -> dict[str, Any]:
    return {
        "selectionMode": "RECENT_DOCUMENT_ACTIVITY",
        **recent_window(today),
        "dateFieldLabel": "创建日期",
        "documentType": "ALL",
        "jurisdiction": "全部管辖单位(含派出所)",
        "brigadeScope": "ALL",
        "timezone": "Asia/Shanghai",
    }


def effective_result(product: dict[str, Any]) -> str:
    if product.get("method") == "SAMPLING" and product.get("reinspectionApplied") == "YES":
        result = product.get("reinspectionResult")
        return result if result in {"QUALIFIED", "UNQUALIFIED"} else "PENDING"
    result = product.get("result")
    return result if result in {"QUALIFIED", "UNQUALIFIED", "PENDING"} else "UNKNOWN"


def qualification(manifest: dict[str, Any]) -> dict[str, Any]:
    initial = manifest.get("initialInspection") or {}
    products = initial.get("products") or []
    outcomes = [effective_result(p) for p in products if isinstance(p, dict)]
    result = (
        "UNQUALIFIED"
        if "UNQUALIFIED" in outcomes
        else "QUALIFIED"
        if outcomes and all(r == "QUALIFIED" for r in outcomes)
        else "UNKNOWN"
    )
    raw_date = initial.get("inspectionDate")
    try:
        inspection_date = date.fromisoformat(raw_date[:10]).isoformat() if raw_date else None
    except (ValueError, TypeError):
        inspection_date = None
    return {
        "initialResult": result,
        "initialInspectionDate": inspection_date,
        "statisticsYear": int(inspection_date[:4]) if inspection_date else None,
        "reinspectionPending": "PENDING" in outcomes,
    }


def snapshot_qualification(snapshot: dict[str, Any]) -> dict[str, Any]:
    """CaseImportState stores business values with their provenance, not flat fields."""
    initial = next(
        (item for item in snapshot.get("inspections", []) if item.get("stage") == "INITIAL_CHECK"),
        {},
    )

    def values(entity: dict[str, Any]) -> dict[str, Any]:
        return {
            key: field.get("value")
            for key, field in (entity.get("fields") or {}).items()
            if isinstance(field, dict)
        }

    return qualification(
        {
            "initialInspection": {
                **values(initial),
                "products": [values(product) for product in initial.get("products", [])],
            }
        }
    )


def source_qualification(fields: dict[str, Any], stages: list[str]) -> dict[str, Any] | None:
    if isinstance(fields.get("initialInspection"), dict):
        return qualification(fields)
    # The source UI's exact initial-only, closed onsite checks are structured
    # evidence. Tags or a document title containing "合格" alone are not.
    products = fields.get("检查产品信息")
    if stages != ["INITIAL"] or fields.get("检查情况") != "已结案[合格]":
        return None
    if (
        not isinstance(products, list)
        or not products
        or not all(
            isinstance(p, dict)
            and p.get("产品质量现场检查情况") == "未发现不合格现象"
            and p.get("市场准入检查情况") == "未发现不合格现象"
            and p.get("检查结果") in (None, "", "合格")
            for p in products
        )
    ):
        return None
    return qualification(
        {
            "initialInspection": {
                "inspectionDate": fields.get("初查日期") or fields.get("实际检查日期"),
                "products": [{"method": "ONSITE", "result": "QUALIFIED"} for _ in products],
            }
        }
    )


def classify_saved_details(layout: Any, *, apply: bool = False) -> list[dict[str, Any]]:
    changes = []
    for project, record in ws.load_waterline(layout)["cases"].items():
        if record.get("source", {}).get("projectIdentitySource") != "DETAIL":
            continue
        _, manifest, _, _ = case_inputs(layout, project, record)
        if manifest or record.get("source", {}).get("changePending"):
            continue
        path = layout.pending_case_dir(project) / "source-evidence.json"
        evidence = _read(path, layout.root)
        if evidence.get("projectNo") != project:
            continue
        candidates = [
            source_qualification(item.get("fields", {}), item.get("projectInspectionStages", []))
            for item in evidence.get("records", {}).values()
            if item.get("projectNo") == project
        ]
        if (
            not candidates
            or any(c is None or c["initialResult"] == "UNKNOWN" for c in candidates)
            or any(c != candidates[0] for c in candidates)
        ):
            continue
        classification = {
            **candidates[0],
            "evidencePath": str(path.relative_to(layout.root)),
            "evidenceSha256": "sha256:" + hashlib.sha256(path.read_bytes()).hexdigest(),
            "observedAt": evidence.get("updatedAt"),
        }
        if record.get("source", {}).get("classification") == classification:
            continue
        if apply:
            ws.upsert_case(
                layout,
                project,
                source={
                    "classification": classification,
                    "indexOnly": classification["initialResult"] == "QUALIFIED",
                },
            )
        changes.append(
            {
                "projectNo": project,
                "initialResult": classification["initialResult"],
                "statisticsYear": classification["statisticsYear"],
                "applied": apply,
            }
        )
    return changes


def _read(path: Path, root: Path) -> dict[str, Any]:
    if not path.exists():
        return {}
    return ws._read_progress_json(path, root, "总账引用证据")


def case_inputs(layout: Any, project: str, record: dict[str, Any]) -> tuple[Path, dict, dict, list]:
    """Only follow the exact active directory or the ledger's explicit archived generation."""
    work = layout.work_case_dir(project)
    issues: list[str] = []
    if not work.is_dir():
        archive = record.get("archive") or {}
        candidate = archive.get("workspacePath")
        if candidate:
            candidate = Path(candidate)
            suffix = candidate.name.removeprefix(project + "-")
            if (
                candidate.parent != layout.root / "工作区" / "历史工作区"
                or not candidate.name.startswith(project + "-")
                or not ws.ARCHIVE_GENERATION.fullmatch(suffix)
            ):
                return work, {}, {}, ["ARCHIVE_REFERENCE_INVALID"]
            work = candidate
    try:
        manifest = _read(work / "manifest.json", layout.root)
        state = _read(work / "upload-state.json", layout.root)
        if manifest and manifest.get("case", {}).get("projectNo") != project:
            raise ws.WorkspaceStateError("清单项目身份不符")
        if state and (state.get("stateVersion") != 6 or state.get("projectNo") != project):
            raise ws.WorkspaceStateError("上传断点项目身份不符")
        if (
            manifest
            and state
            and state.get("manifestSha256")
            != ("sha256:" + hashlib.sha256((work / "manifest.json").read_bytes()).hexdigest())
        ):
            raise ws.WorkspaceStateError("清单与断点摘要不符")
        return work, manifest, state, issues
    except (ws.WorkspaceStateError, OSError, ValueError):
        return work, {}, {}, ["LOCAL_EVIDENCE_UNTRUSTED"]


def _content_verified(work: Path, layout: Any, manifest: dict, state: dict) -> bool:
    proof = _read(work / "content-verification.json", layout.root)
    if (
        not proof
        or proof.get("schemaVersion") != "ContentVerificationV1"
        or not proof.get("completedAt")
    ):
        return False
    if (
        proof.get("manifestSha256") != state.get("manifestSha256")
        or proof.get("caseId") != state.get("caseId")
        or proof.get("origin") != state.get("origin")
    ):
        return False
    files = manifest.get("files") or []
    observed = proof.get("files") or {}
    return (
        bool(files)
        and len(observed) == len(files)
        and all(
            isinstance(observed.get(f["clientRef"]), dict)
            and observed[f["clientRef"]].get("sha256") == f.get("sha256")
            and observed[f["clientRef"]].get("verifiedAt")
            and observed[f["clientRef"]].get("fileId")
            and observed[f["clientRef"]].get("verifiedAgainstManifestSha256")
            and type(observed[f["clientRef"]].get("contentGeneration")) is int
            and observed[f["clientRef"]]["contentGeneration"] >= 1
            for f in files
        )
    )


def describe(layout: Any, project: str, record: dict[str, Any]) -> dict[str, Any]:
    work, manifest, state, issues = case_inputs(layout, project, record)
    source = record.get("source") or {}
    archive = record.get("archive") or {}
    summary = state.get("finalizeSummary") or {}
    registered = bool(
        state.get("caseId")
        and state.get("status") in {"FINALIZED_UNVERIFIED", "FINALIZED_WITH_CONFLICTS", "VERIFIED"}
        and summary.get("created") is True
    )
    binding = record.get("systemObservation") or {}
    registered = registered or bool(
        state.get("caseId")
        and state.get("status") == "VERIFIED"
        and state.get("verification", {}).get("caseId") == state["caseId"]
    )
    registered = registered or bool(binding.get("exists") is True and binding.get("caseId"))
    if binding.get("exists") is False and binding.get("observedAt"):
        registered = False
    try:
        body_verified = _content_verified(work, layout, manifest, state) if manifest else False
    except (ws.WorkspaceStateError, OSError, KeyError, TypeError):
        body_verified = False
        issues.append("CONTENT_RECEIPT_UNTRUSTED")
    system_changed = bool(state.get("caseId") and binding.get("exists") is False)
    if system_changed:
        body_verified = False
    if body_verified and isinstance(binding.get("fileSnapshot"), dict):
        proof = _read(work / "content-verification.json", layout.root)
        observed = binding["fileSnapshot"]
        verified = {item["fileId"]: item for item in proof["files"].values()}
        body_current = set(observed) == set(verified) and all(
            item.get("sha256") == verified[file_id].get("sha256")
            and item.get("contentGeneration") == verified[file_id].get("contentGeneration")
            and str(item.get("sizeBytes")) == str(verified[file_id].get("sizeBytes"))
            for file_id, item in observed.items()
        )
        if not body_current:
            body_verified = False
            system_changed = True
            issues.append("SYSTEM_CONTENT_CHANGED_SINCE_VERIFICATION")
    archived = bool(
        archive.get("workspacePath")
        and archive.get("verificationRecord")
        and work.is_dir()
        and state.get("status") == "VERIFIED"
    )
    if archived:
        try:
            receipt = _read(Path(archive["verificationRecord"]), layout.root)
            archived = bool(
                receipt.get("recordVersion") == 1
                and receipt.get("projectNo") == project
                and receipt.get("manifestSha256") == state.get("manifestSha256")
                and receipt.get("archivedWorkspace") == str(work)
            )
        except (OSError, ws.WorkspaceStateError):
            archived = False
            issues.append("ARCHIVE_RECEIPT_UNTRUSTED")
    # Historical completion is retained, but never upgraded to deep-content proof.
    historical_complete = record.get("state") == "COMPLETED" and archived
    projection = state.get("filesProjection") or []
    projected = {f.get("clientRef") for f in projection if isinstance(f, dict)} - {None}
    received = set(state.get("uploadedFileRefs") or []) & projected
    stages = {
        "sourceRegistered": bool(
            source.get("projectIdentitySource") == "DETAIL" or source.get("rwid")
        ),
        "materialsCollected": source.get("status") in {"PACKAGE_READY", "PACKAGE_RECEIVED"}
        or bool(manifest.get("packageSha256") and state),
        "systemRegistered": registered,
        "bodyVerified": body_verified,
        "archived": archived,
    }
    classification = qualification(manifest)
    if (
        binding.get("exists")
        and binding.get("snapshotDigest")
        and isinstance(binding.get("qualification"), dict)
    ):
        classification = {k: binding["qualification"].get(k) for k in classification}
    if (
        not manifest
        and classification["initialResult"] == "UNKNOWN"
        and source.get("classification")
    ):
        candidate = source["classification"]
        try:
            evidence_path = layout.root / candidate["evidencePath"]
            evidence = _read(evidence_path, layout.root)
            if evidence.get("projectNo") == project and candidate["evidenceSha256"] == (
                "sha256:" + hashlib.sha256(evidence_path.read_bytes()).hexdigest()
            ):
                classification = {k: candidate.get(k) for k in classification}
        except (OSError, KeyError, ws.WorkspaceStateError):
            issues.append("CLASSIFICATION_EVIDENCE_UNTRUSTED")
    if source.get("changePending") and source.get("pendingClassification"):
        candidate = source["pendingClassification"]
        try:
            evidence_path = layout.root / candidate["evidencePath"]
            evidence = _read(evidence_path, layout.root)
            if evidence.get("projectNo") != project or candidate["evidenceSha256"] != (
                "sha256:" + hashlib.sha256(evidence_path.read_bytes()).hexdigest()
            ):
                raise ws.WorkspaceStateError("来源变化分类证据不一致")
            classification = {k: candidate.get(k) for k in classification}
        except (OSError, KeyError, ws.WorkspaceStateError):
            issues.append("PENDING_CLASSIFICATION_EVIDENCE_UNTRUSTED")
    source_documents = {
        (rwid, fingerprint)
        for observations in (
            source.get("observationsByRwid"),
            source.get("pendingObservationsByRwid"),
        )
        for rwid, obs in (observations or {}).items()
        for fingerprint in obs.get("documentFingerprints", [])
    }
    return {
        "projectNo": project,
        "unitName": record.get("unitName"),
        **classification,
        "stages": stages,
        "complete": all(stages.values()) and not source.get("changePending") and not system_changed,
        "historicalComplete": historical_complete,
        "activePending": system_changed
        or bool(source.get("changePending"))
        or (
            not historical_complete
            and not (
                source.get("indexOnly") is True and classification["initialResult"] == "QUALIFIED"
            )
        ),
        "evidenceLevel": "DEEP_CONTENT"
        if body_verified
        else "LEGACY_UNSPECIFIED"
        if historical_complete
        else "PENDING",
        "sourceChanged": bool(source.get("changePending")),
        "systemChanged": system_changed,
        "sourceObservedAt": source.get("lastObservedAt") or source.get("capturedAt"),
        "sourceDocumentDate": source.get("latestDocumentCreatedAt"),
        "sourceDocumentCount": len(source_documents),
        "workflowUpdatedAt": record.get("workflowUpdatedAt"),
        "files": {"expected": len(projected), "received": len(received)},
        "waiting": record.get("nasVerification", {}).get("status"),
        "issues": issues,
    }


def ledger_view(layout: Any, *, view: str = "all", batch_id: str | None = None) -> dict[str, Any]:
    ledger = ws.load_waterline(layout)
    records = ledger["cases"]
    if batch_id:
        # Validate the batch exists and is formal; its RWIDs are not project identities.
        ws._formal_capture_batch(layout, batch_id)
        records = {
            p: r
            for p, r in records.items()
            if r.get("source", {}).get("batchId") == batch_id
            or batch_id in r.get("source", {}).get("batchIds", [])
        }
    rows = [describe(layout, p, r) for p, r in sorted(records.items())]
    known = len(rows)
    if view == "unqualified":
        rows = [r for r in rows if r["initialResult"] == "UNQUALIFIED"]
    elif view == "unknown":
        rows = [r for r in rows if r["initialResult"] == "UNKNOWN"]
    elif view == "unfinished":
        rows = [r for r in rows if not r["complete"]]
    elif view != "all":
        raise ws.WorkspaceStateError("总账视图无效")
    stage_names = (
        "sourceRegistered",
        "materialsCollected",
        "systemRegistered",
        "bodyVerified",
        "archived",
    )
    return {
        "schemaVersion": "CaseLedgerViewV1",
        "scope": "KNOWN_CASES",
        "view": view,
        "batchId": batch_id,
        "generatedAt": ws._utc_now(),
        "identityKey": "projectNo",
        "coverage": {
            "knownCases": known,
            "historicalSourceComplete": False,
            "note": "仅已知案卷索引；未证明来源系统全部历史覆盖",
        },
        "counts": {
            "cases": len(rows),
            **{key: sum(r["stages"][key] for r in rows) for key in stage_names},
            "knownSourceDocuments": sum(r["sourceDocumentCount"] for r in rows),
            "complete": sum(r["complete"] for r in rows),
            "historicalComplete": sum(r["historicalComplete"] for r in rows),
            "activePending": sum(r["activePending"] for r in rows),
            "legacyEvidenceIncomplete": sum(
                r["historicalComplete"] and not r["stages"]["bodyVerified"] for r in rows
            ),
            "expectedFiles": sum(r["files"]["expected"] for r in rows),
            "receivedFiles": sum(r["files"]["received"] for r in rows),
        },
        "qualification": dict(Counter(r["initialResult"] for r in rows)),
        "cases": rows,
    }


def render_ledger_report(value: dict[str, Any], *, include_all: bool = False) -> str:
    counts = value["counts"]
    lines = [
        "# 案卷总水位",
        "",
        "所有视图共用项目编号唯一总账；历史来源覆盖未知。",
        "本报告只读本地证据，不等同于本次重新访问服务器核验。",
        "",
        f"本视图 {counts['cases']} 案；材料已采集 {counts['materialsCollected']}；"
        f"系统已登记 {counts['systemRegistered']}；正文已核验 {counts['bodyVerified']}；"
        f"归档 {counts['archived']}；全部完成 {counts['complete']}。",
        f"已知来源文书 {counts['knownSourceDocuments']} 份；"
        f"来源已登记 {counts['sourceRegistered']} 案；"
        f"活动待办 {counts['activePending']} 案。历史完成 {counts['historicalComplete']} 案中，"
        f"{counts['legacyEvidenceIncomplete']} 案缺少新式逐份正文回执，"
        "保留原完成记录，不默认重做。",
        "",
        "|项目编号|初查结果|初查日期|来源登记|材料采集|系统登记|正文核验|归档|当前处理|",
        "|---|---|---|---|---|---|---|---|---|",
    ]
    report_rows = (
        value["cases"]
        if include_all
        else [row for row in value["cases"] if row["activePending"] or row["issues"]]
    )
    if not include_all:
        lines.insert(-2, "以下仅列活动待办或明确异常；历史明细见总账，--all-cases 可显式展开。")
        lines.insert(-2, "")
    for row in report_rows:
        s = row["stages"]
        flags = [
            "是" if s[key] else "未确认"
            for key in (
                "sourceRegistered",
                "materialsCollected",
                "systemRegistered",
                "bodyVerified",
                "archived",
            )
        ]
        lines.append(
            "|"
            + "|".join(
                [
                    row["projectNo"],
                    {"UNQUALIFIED": "不合格", "QUALIFIED": "合格", "UNKNOWN": "待确认"}.get(
                        row["initialResult"], "待确认"
                    ),
                    row["initialInspectionDate"] or "未确认",
                    *flags,
                    "有新增文书"
                    if row["sourceChanged"]
                    else "待处理"
                    if row["activePending"]
                    else "历史证据复用"
                    if row["historicalComplete"]
                    else "仅建索引",
                ]
            )
            + "|"
        )
    return "\n".join(lines) + "\n"


def add_ledger_sheets(workbook: Any, layout: Any) -> None:
    """Export four views of one evidence snapshot; the workbook is never a second ledger."""
    from openpyxl.styles import Alignment, Font, PatternFill
    from openpyxl.utils import get_column_letter

    snapshot = ledger_view(layout)
    all_rows = snapshot["cases"]
    groups = [
        ("所有案卷", all_rows),
        ("不合格案卷", [r for r in all_rows if r["initialResult"] == "UNQUALIFIED"]),
        ("待确认案卷", [r for r in all_rows if r["initialResult"] == "UNKNOWN"]),
        ("未完成案卷", [r for r in all_rows if not r["complete"]]),
    ]
    headers = [
        "项目编号",
        "单位名称",
        "初查结果",
        "初查日期",
        "统计年份",
        "来源已登记",
        "材料已采集",
        "系统已登记",
        "正文已核验",
        "归档已完成",
        "活动待办",
        "来源有变化",
        "最近来源核对",
        "本地处理时间",
        "证据说明",
    ]
    stage_names = (
        "sourceRegistered",
        "materialsCollected",
        "systemRegistered",
        "bodyVerified",
        "archived",
    )
    for index, (name, rows) in enumerate(groups):
        sheet = workbook.create_sheet(name, index)
        sheet.append(headers)
        for row in rows:
            sheet.append(
                [
                    row["projectNo"],
                    row["unitName"] or "",
                    {"QUALIFIED": "合格", "UNQUALIFIED": "不合格", "UNKNOWN": "待确认"}.get(
                        row["initialResult"], "待确认"
                    ),
                    row["initialInspectionDate"] or "",
                    row["statisticsYear"],
                    *("是" if row["stages"][key] else "未确认" for key in stage_names),
                    "是" if row["activePending"] else "否",
                    "是" if row["sourceChanged"] else "否",
                    row["sourceObservedAt"] or "",
                    row["workflowUpdatedAt"] or "",
                    "历史完成，缺少新版正文回执，默认复用"
                    if row["evidenceLevel"] == "LEGACY_UNSPECIFIED"
                    else "逐份正文回执已核对"
                    if row["evidenceLevel"] == "DEEP_CONTENT"
                    else "待处理",
                ]
            )
        sheet.freeze_panes = "A2"
        sheet.auto_filter.ref = sheet.dimensions
        sheet.sheet_view.showGridLines = False
        for cells in sheet.iter_rows():
            for cell in cells:
                if isinstance(cell.value, str):
                    cell.data_type = "s"
                cell.font = Font(
                    name="Microsoft YaHei",
                    size=10,
                    bold=cell.row == 1,
                    color="FFFFFF" if cell.row == 1 else "1F1F1F",
                )
                cell.alignment = Alignment(vertical="center", wrap_text=True)
                if cell.row == 1 or cell.row % 2 == 0:
                    cell.fill = PatternFill(
                        "solid", fgColor="17365D" if cell.row == 1 else "EAF2F8"
                    )
        for column, width in enumerate(
            [24, 34, 12, 15, 10, 13, 13, 13, 13, 13, 12, 12, 25, 25, 42], 1
        ):
            sheet.column_dimensions[get_column_letter(column)].width = width
        sheet.row_dimensions[1].height = 32
    workbook.active = 0
    info = workbook["使用说明"]
    info.append(
        ["覆盖范围", "仅已知案卷；不代表来源系统全部历史。四种视图按项目编号共用同一总账。"]
    )
    info.append(["历史证据", "未完成视图包含证据不足的历史完成案；活动待办为否时不默认重新采集。"])


def scan_plan(layout: Any, batch_id: str | None = None) -> dict[str, Any]:
    ledger = ws.load_waterline(layout)
    pending = []
    for project, record in ledger["cases"].items():
        row = describe(layout, project, record)
        if row["activePending"]:
            pending.append(
                {
                    "projectNo": project,
                    "initialResult": row["initialResult"],
                    "changeEvidence": list(
                        record.get("source", {}).get("changeEvidenceByFingerprint", {}).values()
                    )
                    if row["sourceChanged"]
                    else [],
                    "action": "RECONCILE"
                    if row["stages"]["systemRegistered"] or row["systemChanged"]
                    else "COLLECT"
                    if row["initialResult"] == "UNQUALIFIED"
                    else "CLASSIFY"
                    if row["initialResult"] == "UNKNOWN"
                    else "INDEX_ONLY",
                }
            )
    pending.sort(key=lambda x: (x["initialResult"] != "UNQUALIFIED", x["projectNo"]))
    actions = []
    if batch_id:
        if __package__:
            from .source_intake import document_fingerprint
        else:
            from source_intake import document_fingerprint
        _, _, capture = ws._formal_capture_batch(layout, batch_id)
        if capture.get("stableRounds", 0) < 2:
            raise ws.WorkspaceStateError("窗口清单尚未连续两轮稳定")
        by_rwid = {}
        for project, record in ledger["cases"].items():
            source = record.get("source") or {}
            for rwid in [source.get("rwid"), *(source.get("observationsByRwid") or {})]:
                if rwid:
                    by_rwid.setdefault(rwid, []).append(project)
        for rwid, observation in capture.get("records", {}).items():
            projects = set(by_rwid.get(rwid, []))
            project = next(iter(projects)) if len(projects) == 1 else None
            record = ledger["cases"].get(project, {})
            source = record.get("source") or {}
            previous = source.get("observationsByRwid", {}).get(rwid, {})
            incoming = {document_fingerprint(a) for a in observation.get("sourceAppearances", [])}
            same = bool(incoming and incoming <= set(previous.get("documentFingerprints") or []))
            unchanged_complete = (
                same
                and not source.get("changePending")
                and not describe(layout, project, record)["systemChanged"]
                and (
                    describe(layout, project, record)["historicalComplete"]
                    or (
                        source.get("indexOnly") is True
                        and describe(layout, project, record)["initialResult"] == "QUALIFIED"
                    )
                )
            )
            actions.append(
                {
                    "rwid": rwid,
                    "projectNo": project,
                    "action": "IDENTITY_CONFLICT"
                    if len(projects) > 1
                    else "SKIP_UNCHANGED"
                    if unchanged_complete
                    else "RESUME"
                    if same
                    else "READ_DETAIL",
                    "sourceChanged": bool(previous and not same),
                }
            )
    return {
        "schemaVersion": "IncrementalCasePlanV1",
        "readOnly": True,
        "filters": scan_filters(),
        "pendingOutsideWindowRetained": True,
        "pending": pending,
        "sourceActions": actions,
    }


def index_history(layout: Any, *, apply: bool = False) -> dict[str, Any]:
    """Reuse formal, locally saved lists. No browser, package read, or OCR."""
    if __package__:
        from .source_intake import document_fingerprint
    else:
        from source_intake import document_fingerprint
    results = []
    missing_identity = set()
    for directory in sorted(layout.capture_batches.iterdir()):
        if not directory.is_dir():
            continue
        try:
            _, _, capture = ws._formal_capture_batch(layout, directory.name)
        except ws.WorkspaceStateError:
            continue
        for rwid, item in capture.get("records", {}).items():
            detail = item.get("detail") or {}
            project = detail.get("projectNo")
            fields = detail.get("fields") or {}
            if (
                not project
                or fields.get("projectNo") != project
                or not ws.PROJECT_NO.fullmatch(project)
            ):
                missing_identity.add(rwid)
                continue
            observation = {
                "fingerprint": item.get("sourceRecordFingerprint"),
                "documentFingerprints": sorted(
                    {document_fingerprint(a) for a in item.get("sourceAppearances", [])}
                ),
                "observedAt": detail.get("capturedAt"),
            }
            record = ws.load_waterline(layout)["cases"].get(project, {})
            prior = record.get("source", {}).get("observationsByRwid", {}).get(rwid, {})
            new_docs = set(observation["documentFingerprints"]) - set(
                prior.get("documentFingerprints", [])
            )
            if project in ws.load_waterline(layout)["cases"] and not new_docs:
                continue
            if apply:
                updates = {
                    "observationsByRwid": {rwid: observation},
                    "batchIds": sorted(
                        set(record.get("source", {}).get("batchIds", [])) | {directory.name}
                    ),
                }
                if not record:
                    updates.update(projectIdentitySource="DETAIL", rwid=rwid, status="DISCOVERED")
                ws.upsert_case(
                    layout,
                    project,
                    source=updates,
                    **({"unitName": fields.get("unitName")} if not record else {}),
                )
            results.append(
                {"projectNo": project, "newKnownDocuments": len(new_docs), "applied": apply}
            )
    return {
        "readOnly": not apply,
        "historicalSourceComplete": False,
        "sourceIdentitiesWithoutProject": len(missing_identity),
        "changes": results,
        "classificationChanges": classify_saved_details(layout, apply=apply),
    }
