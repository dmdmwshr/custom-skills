import hashlib
from pathlib import Path
import tempfile
import unittest

from docx import Document
from docx.oxml.ns import qn

from review_docx import build, source_fields


class ReviewDocumentTests(unittest.TestCase):
    def setUp(self):
        self.tmp = tempfile.TemporaryDirectory()
        self.addCleanup(self.tmp.cleanup)
        self.root = Path(self.tmp.name)
        self.source = self.root / "source.docx"
        d = Document()
        d.add_paragraph("无锡市消防救援局XXX项目技术需求分析报告论证会专家评审意见")
        t = d.add_table(rows=2, cols=2)
        t.cell(0, 0).text = "某专家"
        t.cell(0, 1).text = "某单位研究所所长/高工"
        t.cell(1, 0).text = "XXX"
        t.cell(1, 1).text = "基层代表：XXX"
        d.add_paragraph("2026年9月XX日下午组织召开评审会，形成意见如下：")
        d.add_paragraph("1、明确起算条件。")
        d.add_paragraph("2、建议核对考核口径。")
        d.add_paragraph("3、补充验收依据。")
        d.add_paragraph("用户已确定的评审结论。")
        d.add_paragraph("专家签字：")
        d.save(self.source)

    def test_normalization_preserves_body_and_source(self):
        digest = hashlib.sha256(self.source.read_bytes()).digest()
        result = self.root / "out.docx"
        original_body = [p.text for p in Document(self.source).paragraphs][1:]
        build(self.source, result)
        d = Document(result)
        self.assertFalse(d.tables)
        self.assertEqual([p.text for p in d.paragraphs][-len(original_body):], original_body)
        self.assertEqual(hashlib.sha256(self.source.read_bytes()).digest(), digest)
        self.assertIn("某专家\t某单位研究所所长/高工", [p.text for p in d.paragraphs])
        self.assertIn("XXX\tXXXXXXXXXXXXX", [p.text for p in d.paragraphs])
        self.assertTrue(d.sections[0].different_first_page_header_footer)
        self.assertEqual(d.sections[0].first_page_footer.paragraphs[0].text, "2026年9月XX日")

    def test_prefix_is_opt_in_and_idempotent(self):
        first, second = self.root / "one.docx", self.root / "two.docx"
        build(self.source, first, prefix_suggestions=True)
        build(first, second, prefix_suggestions=True)
        lines = [p.text for p in Document(second).paragraphs]
        self.assertIn("1、建议明确起算条件。", lines)
        self.assertIn("2、建议核对考核口径。", lines)
        self.assertIn("3、建议补充验收依据。", lines)
        self.assertFalse(any("建议建议" in p for p in lines))
        self.assertIn("用户已确定的评审结论。", lines)

    def test_blank_template_does_not_retain_people_or_approve_project(self):
        result = self.root / "blank.docx"
        build(self.source, result, blank=True)
        d = Document(result)
        text = "\n".join(p.text for p in d.paragraphs)
        self.assertNotIn("某专家", text)
        self.assertNotIn("用户已确定的评审结论", text)
        self.assertIn("评审结论：XXX。", text)
        self.assertEqual(text.count("、建议XXX"), 3)
        for r in d._element.iter(qn("w:r")):
            pr = r.find(qn("w:rPr"))
            if pr is not None and pr.find(qn("w:highlight")) is not None:
                self.assertEqual(set("".join(t.text or "" for t in r.iter(qn("w:t")))), {"X"})

    def test_existing_output_and_wrong_opinion_count_are_rejected(self):
        with self.assertRaises(ValueError):
            build(self.source, self.source)
        content = {"title": "测试", "experts": [["XXX", "XXX"]], "date": "XXXX年XX月XX日",
                   "introduction": "XXXX年XX月XX日组织召开会议，形成意见：", "opinions": ["建议核对。"], "conclusion": "XXX"}
        target = self.root / "invalid.docx"
        with self.assertRaises(ValueError):
            build(self.source, target, content=content)
        self.assertFalse(target.exists())

    def test_reviewed_content_replaces_template_examples(self):
        content = {"title": "项目技术需求分析报告论证会专家评审意见", "experts": [["某专家", "某单位研究所所长/高工"], ["XXX", "XXXXXXXXXXXXX"]],
                   "date": "2026年9月XX日", "introduction": "2026年9月XX日组织召开会议，形成意见如下：",
                   "opinions": ["建议明确服务范围。", "建议统一计费口径。", "建议补充验收依据。"], "conclusion": "XXX"}
        result = self.root / "generated.docx"
        build(self.source, result, content=content)
        d = Document(result)
        text = "\n".join(p.text for p in d.paragraphs)
        self.assertEqual(d.paragraphs[0].text, content["title"])
        self.assertIn("2、建议统一计费口径。", text)
        self.assertNotIn("用户已确定的评审结论", text)
        self.assertFalse(d.tables)


if __name__ == "__main__":
    unittest.main()
