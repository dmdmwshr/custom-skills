"""Read-only labor-source comparison. Never print full personal identifiers.

Selection JSON contains roster path and experts [{name, history, sheet?}].
No output files are written; exit 2 means stop and confirm source conflicts.
"""
from __future__ import annotations
import argparse
import json
from pathlib import Path
import re

ALIASES = {
    '姓名': ['姓名'], '身份证号': ['身份证号', '身份证号码'],
    '工作单位': ['工作单位', '单位'],
    '开户行网点': ['开户行网点', '开户行', '开户银行'],
    '银行卡号': ['银行卡号', '银行账号'],
    '手机': ['手机', '手机号', '手机号码', '联系电话'],
}
IDENTIFIERS = {'身份证号', '银行卡号'}


def compact(value):
    return re.sub(r'\s+', '', str(value or ''))


def read_rows(path, sheet=None):
    path = Path(path)
    if path.suffix.lower() == '.xls':
        import xlrd
        book = xlrd.open_workbook(str(path))
        try:
            sh = book.sheet_by_name(sheet) if sheet else book.sheet_by_index(0)
            return [sh.row_values(i) for i in range(sh.nrows)]
        finally:
            book.release_resources()
    if path.suffix.lower() in {'.xlsx', '.xlsm'}:
        from openpyxl import load_workbook
        book = load_workbook(path, read_only=True, data_only=True)
        try:
            sh = book[sheet] if sheet else book.worksheets[0]
            return list(sh.iter_rows(values_only=True))
        finally:
            book.close()
    raise ValueError('Only XLS/XLSX/XLSM sources are supported')


def find_record(rows, name):
    header = next((i for i, row in enumerate(rows[:10]) if '姓名' in [compact(v) for v in row]
                   and any(compact(v) in ALIASES['身份证号'] for v in row)), None)
    if header is None:
        raise ValueError('Missing recognizable name/identity header')
    labels = [compact(v) for v in rows[header]]
    columns = {}
    for field, aliases in ALIASES.items():
        matches = [i for i, label in enumerate(labels) if label in aliases]
        if len(matches) != 1:
            raise ValueError('Missing or duplicate header: ' + field)
        columns[field] = matches[0]
    matches = [(i+1, row) for i, row in enumerate(rows[header+1:], header+1)
               if len(row) > columns['姓名'] and compact(row[columns['姓名']]) == name]
    if len(matches) != 1:
        raise ValueError('Expected one matching person; found ' + str(len(matches)))
    number, row = matches[0]
    return number, {field: row[index] if index < len(row) else None for field,index in columns.items()}


def normalize(value, field):
    if value is None or compact(value) == '':
        raise ValueError('missing_value')
    if field in IDENTIFIERS and not isinstance(value, str):
        raise ValueError('identifier_not_text_precision_unverified')
    if isinstance(value, (float,int)) and not isinstance(value, bool):
        if int(value) != value:
            raise ValueError('non_integer_number')
        value = str(int(value))
    result = compact(value).upper() if field == '身份证号' else compact(value)
    if field == '身份证号' and not re.fullmatch(r'(?:\d{15}|\d{17}[\dX])',result):
        raise ValueError('invalid_identifier_format')
    if field == '银行卡号' and not re.fullmatch(r'\d{12,19}',result):
        raise ValueError('invalid_account_format')
    if field == '手机' and not re.fullmatch(r'\d{11}',result):
        raise ValueError('invalid_phone_format')
    return result


def redact(value, field):
    if field in IDENTIFIERS or field == '手机':
        return '***' + value[-4:]
    return value


def audit(selection):
    experts = selection['experts']
    names = [compact(e['name']) for e in experts]
    if not names or any(not n for n in names) or len(names) != len(set(names)):
        raise ValueError('Selection must contain distinct confirmed expert names')
    roster = read_rows(selection['roster'],selection.get('roster_sheet','专家库'))
    results=[]
    for expert,name in zip(experts,names):
        result={'name':name,'history':expert['history'],'conflicts':[]}
        try:
            old_number,old=find_record(read_rows(expert['history'],expert.get('sheet')),name)
            new_number,new=find_record(roster,name)
            result.update(history_row=old_number,roster_row=new_number)
            for field in ALIASES:
                if field == '姓名':continue
                try:
                    a,b=normalize(old[field],field),normalize(new[field],field)
                except ValueError as error:
                    result['conflicts'].append({'field':field,'reason':str(error)})
                    continue
                if a != b:
                    result['conflicts'].append({'field':field,'reason':'mismatch',
                                                'history':redact(a,field),'roster':redact(b,field)})
        except (ValueError,KeyError,IndexError) as error:
            result['conflicts'].append({'field':'record','reason':str(error)})
        results.append(result)
    blocked=any(r['conflicts'] for r in results)
    return {'status':'needs_confirmation' if blocked else 'matched', 'experts':results,
            'remaining_checks':['岗位职称及简称语义核验','本次人数金额日期确认','正确修订版及历史签字边界核验']}


def main():
    parser=argparse.ArgumentParser(description=__doc__)
    parser.add_argument('--selection',required=True,help='Task-local JSON: paths/names only, no identifiers')
    args=parser.parse_args()
    try:
        report=audit(json.loads(Path(args.selection).read_text(encoding='utf-8-sig')))
    except (OSError,ValueError,KeyError,ImportError) as error:
        print(json.dumps({'status':'error','reason':str(error)},ensure_ascii=False))
        return 2
    print(json.dumps(report,ensure_ascii=False,indent=2))
    return 0 if report['status']=='matched' else 2


if __name__ == '__main__':
    raise SystemExit(main())
