import hashlib
import json
from pathlib import Path
import tempfile
import unittest
from openpyxl import Workbook,load_workbook
from audit_labor_sources import audit


class LaborAuditTests(unittest.TestCase):
    def setUp(self):
        self.tmp=tempfile.TemporaryDirectory();self.addCleanup(self.tmp.cleanup)
        self.root=Path(self.tmp.name)
        self.roster=self.root/'roster.xlsx';self.history=self.root/'history.xlsx'
        # Deliberately fictional identities/accounts; never copy business PII into tests.
        self.values=['测试专家甲','00000000000000000X','测试单位','测试银行网点','0000000000000000','10000000000']
        for path,title in [(self.roster,'专家库'),(self.history,'Sheet1')]:
            book=Workbook();sheet=book.active;sheet.title=title
            sheet.append(['测试标题'])
            sheet.append(['姓名','身份证号','工作单位','开户行网点','银行卡号','手机号码'])
            sheet.append(self.values);book.save(path)
        self.selection={'roster':str(self.roster),'experts':[{'name':'测试专家甲','history':str(self.history)}]}

    def change(self,cell,value):
        book=load_workbook(self.history);book.active[cell]=value;book.save(self.history);book.close()

    def test_match_preserves_files_and_does_not_expose_identifiers(self):
        self.change('B3','00000000000000000x ')
        self.change('F3',10000000000)
        hashes=[hashlib.sha256(p.read_bytes()).digest() for p in [self.roster,self.history]]
        result=audit(self.selection)
        self.assertEqual(result['status'],'matched')
        self.assertEqual(result['experts'][0]['history_row'],3)
        self.assertEqual(hashes,[hashlib.sha256(p.read_bytes()).digest() for p in [self.roster,self.history]])
        for value in [self.values[i] for i in [1,4,5]]:
            self.assertNotIn(value,json.dumps(result))

    def test_account_conflict_stops_and_redacts(self):
        self.change('E3','0000000000000001')
        result=audit(self.selection)
        self.assertEqual(result['status'],'needs_confirmation')
        conflict=result['experts'][0]['conflicts'][0]
        self.assertEqual(conflict['field'],'银行卡号')
        self.assertEqual(conflict['history'],'***0001')
        self.assertNotIn('0000000000000001',json.dumps(result))

    def test_numeric_long_identifier_is_not_silently_accepted(self):
        self.change('B3',123456789012345678)
        result=audit(self.selection)
        self.assertEqual(result['status'],'needs_confirmation')
        self.assertIn('precision_unverified',result['experts'][0]['conflicts'][0]['reason'])

    def test_missing_and_duplicate_records_block(self):
        self.change('D3',None)
        self.assertEqual(audit(self.selection)['status'],'needs_confirmation')
        book=load_workbook(self.history);book.active.append(self.values);book.save(self.history);book.close()
        self.assertEqual(audit(self.selection)['status'],'needs_confirmation')
        self.selection['experts']*=2
        with self.assertRaises(ValueError):audit(self.selection)


if __name__=='__main__':unittest.main()
